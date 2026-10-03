"""Daily / weekly / monthly reports plus CSV and XLSX export."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from app.models.attendance import (
    ATTENDANCE_CLOSED,
    ATTENDANCE_MISSING,
    ATTENDANCE_OPEN,
    AttendanceRecord,
    AttendanceStatus,
)
from app.models.audit import AuditLog
from app.models.employee import Employee
from app.services.attendance_service import AttendanceService
from app.services.settings_service import SettingsService
from app.utils.time_utils import TimeUtils, WorkWeek

XLSX_AVAILABLE = True
try:  # openpyxl is optional; CSV always works.
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - optional dependency
    XLSX_AVAILABLE = False


@dataclass
class ReportRow:
    values: list = field(default_factory=list)


@dataclass
class ReportResult:
    title: str
    subtitle: str
    headers: list[str]
    rows: list[list]
    summary: list[tuple[str, str]] = field(default_factory=list)
    totals: list[str] = field(default_factory=list)

    @property
    def is_empty(self) -> bool:
        return not self.rows


class ReportService:
    def __init__(self, attendance: AttendanceService, settings: SettingsService) -> None:
        self.attendance = attendance
        self.settings = settings
        self.repos = attendance.repos
        self.clock = settings.time_utils()

    # -- daily ---------------------------------------------------------------
    def daily_report(
        self, work_date: date | None = None, employee_id: int | None = None
    ) -> ReportResult:
        clock = self.clock
        day = work_date or clock.today()
        key = day.strftime("%Y-%m-%d")
        records = self.repos.attendance.list_for_date(key, employee_id)
        now = clock.now()
        now_iso = TimeUtils.to_iso(now)

        headers = ["Employee ID", "Full Name", "Department", "Time In", "Time Out", "Hours", "Status"]
        rows: list[list] = []
        totals_minutes = 0
        working = closed = missing = 0

        for record in records:
            minutes = record.duration_minutes or 0
            if record.is_open:
                minutes = TimeUtils.elapsed_minutes(
                    TimeUtils.parse(record.time_in) or now, now
                )
            totals_minutes += minutes
            if record.status == ATTENDANCE_MISSING:
                missing += 1
            elif record.is_open:
                working += 1
            else:
                closed += 1
            rows.append(
                [
                    record.employee_code,
                    record.full_name,
                    record.department or "-",
                    clock.format_time(record.time_in),
                    clock.format_time(record.time_out),
                    clock.format_duration(minutes),
                    record.status_label,
                ]
            )

        if employee_id is None:
            present = {record.employee_id for record in records}
            for employee in self.repos.employees.list_all(status="active"):
                if employee.employee_id not in present:
                    rows.append(
                        [
                            employee.employee_code,
                            employee.full_name,
                            employee.department or "-",
                            "-",
                            "-",
                            "0h 00m",
                            AttendanceStatus.label(AttendanceStatus.NOT_IN),
                        ]
                    )

        summary = [
            ("Date", clock.format_long_date(day)),
            ("Employees present", str(len(records))),
            ("Currently working", str(working)),
            ("Timed out", str(closed)),
            ("Missing time out", str(missing)),
            ("Total hours", clock.format_duration(totals_minutes)),
        ]
        return ReportResult(
            title=f"Daily Attendance Report - {clock.format_date(day)}",
            subtitle=f"{self.settings.settings.organization_name}  •  {key}",
            headers=headers,
            rows=rows,
            summary=summary,
            totals=["TOTAL", "", "", "", "", clock.format_duration(totals_minutes), f"{len(rows)} row(s)"],
        )

    # -- weekly --------------------------------------------------------------
    def weekly_report(
        self,
        window: WorkWeek | None = None,
        employee_id: int | None = None,
    ) -> ReportResult:
        clock = self.clock
        window = window or clock.week_bounds()
        start = window.start.strftime("%Y-%m-%d")
        end = window.end.strftime("%Y-%m-%d")
        now_iso = TimeUtils.to_iso(clock.now())

        headers = [
            "Employee ID",
            "Full Name",
            "Department",
            "Hours Worked",
            "Weekly Goal",
            "Remaining",
            "Overtime",
            "Days",
            "Status",
        ]
        employees = self.repos.employees.list_all()
        if employee_id is not None:
            employees = [e for e in employees if e.employee_id == employee_id]

        rows: list[list] = []
        total_worked = 0
        total_goal = 0
        for employee in employees:
            worked = self.repos.attendance.minutes_for_employee(
                employee.employee_id, start, end, now_iso
            )
            goal = self.settings.goal_minutes_for(employee.weekly_goal_hours)
            remaining = max(0, goal - worked)
            overtime = max(0, worked - goal) if goal else 0
            totals = self.repos.attendance.daily_totals(employee.employee_id, start, end)
            days = sum(1 for value in totals.values() if int(value.get("minutes") or 0) > 0)
            status = self._weekly_status(worked, goal, employee)
            rows.append(
                [
                    employee.employee_code,
                    employee.full_name,
                    employee.department or "-",
                    clock.format_duration(worked),
                    clock.format_duration(goal),
                    clock.format_duration(remaining),
                    clock.format_duration(overtime, always_sign=True) if overtime else "-",
                    str(days),
                    status,
                ]
            )
            total_worked += worked
            total_goal += goal

        summary = [
            ("Week", clock.week_label()),
            ("Period", f"{clock.format_date(window.start)} - {clock.format_date(window.end)}"),
            ("Employees", str(len(employees))),
            ("Total hours", clock.format_duration(total_worked)),
            ("Total goal", clock.format_duration(total_goal)),
            ("Overtime", clock.format_duration(max(0, total_worked - total_goal), always_sign=True)),
        ]
        return ReportResult(
            title=f"Weekly Attendance Report - {clock.week_label()}",
            subtitle=(
                f"{self.settings.settings.organization_name}  •  "
                f"{clock.format_date(window.start)} to {clock.format_date(window.end)}"
            ),
            headers=headers,
            rows=rows,
            summary=summary,
            totals=["TOTAL", "", "", clock.format_duration(total_worked), clock.format_duration(total_goal),
                    clock.format_duration(max(0, total_goal - total_worked)), "", f"{len(rows)} row(s)", ""],
        )

    def _weekly_status(self, worked: int, goal: int, employee: Employee) -> str:
        if not goal:
            return "No goal set"
        if worked == 0:
            return "Not started"
        if worked >= goal:
            overtime = worked - goal
            return f"Goal reached +{clock_format(overtime)}" if overtime else "Goal reached"
        return "In progress"

    # -- monthly -------------------------------------------------------------
    def monthly_report(
        self,
        window: WorkWeek | None = None,
        employee_id: int | None = None,
    ) -> ReportResult:
        clock = self.clock
        window = window or clock.month_bounds()
        start = window.start.strftime("%Y-%m-%d")
        end = window.end.strftime("%Y-%m-%d")
        now_iso = TimeUtils.to_iso(clock.now())

        headers = ["Date", "Employee ID", "Full Name", "Time In", "Time Out", "Hours", "Status"]
        records = self.repos.attendance.list_for_range(start, end, employee_id)
        records.sort(key=lambda r: (r.work_date, r.full_name))

        rows: list[list] = []
        totals_minutes = 0
        for record in records:
            minutes = record.duration_minutes or 0
            if record.is_open:
                minutes = TimeUtils.elapsed_minutes(
                    TimeUtils.parse(record.time_in) or clock.now(), clock.now()
                )
            totals_minutes += minutes
            rows.append(
                [
                    clock.format_date(record.work_date),
                    record.employee_code,
                    record.full_name,
                    clock.format_time(record.time_in),
                    clock.format_time(record.time_out),
                    clock.format_duration(minutes),
                    record.status_label,
                ]
            )

        summary_rows = self.repos.attendance.month_summary(start, end)
        summary_lines = [
            ("Period", f"{clock.format_date(window.start)} - {clock.format_date(window.end)}"),
            ("Records", str(len(records))),
            ("Employees with records", str(len(summary_rows))),
            ("Total hours", clock.format_duration(totals_minutes)),
        ]
        for entry in summary_rows:
            goal = self.settings.goal_minutes_for(entry["weekly_goal_hours"])
            summary_lines.append(
                (
                    f"  {entry['full_name']}",
                    f"{clock.format_duration(entry['minutes'])} / "
                    f"{clock.format_duration(goal)}"
                    + (f"  (+{clock.format_duration(entry['minutes'] - goal, always_sign=True)})"
                       if goal and entry["minutes"] > goal else ""),
                )
            )

        return ReportResult(
            title=f"Monthly Attendance Report - {window.start.strftime('%B %Y')}",
            subtitle=f"{self.settings.settings.organization_name}  •  {start} to {end}",
            headers=headers,
            rows=rows,
            summary=summary_lines,
            totals=["TOTAL", "", "", "", "", clock.format_duration(totals_minutes), f"{len(rows)} row(s)"],
        )

    # -- custom range --------------------------------------------------------
    def range_report(
        self, start: date, end: date, employee_id: int | None = None
    ) -> ReportResult:
        clock = self.clock
        records = self.repos.attendance.list_for_range(
            start.strftime("%Y-%m-%d"), end.strftime("%Y-%m-%d"), employee_id
        )
        headers = ["Date", "Employee ID", "Full Name", "Department", "Time In", "Time Out", "Hours", "Status"]
        rows: list[list] = []
        totals_minutes = 0
        for record in records:
            minutes = record.duration_minutes or 0
            if record.is_open:
                minutes = TimeUtils.elapsed_minutes(
                    TimeUtils.parse(record.time_in) or clock.now(), clock.now()
                )
            totals_minutes += minutes
            rows.append(
                [
                    clock.format_date(record.work_date),
                    record.employee_code,
                    record.full_name,
                    record.department or "-",
                    clock.format_time(record.time_in),
                    clock.format_time(record.time_out),
                    clock.format_duration(minutes),
                    record.status_label,
                ]
            )
        days = len(clock.date_range(start, end))
        return ReportResult(
            title="Custom Range Attendance Report",
            subtitle=(
                f"{self.settings.settings.organization_name}  •  "
                f"{clock.format_date(start)} to {clock.format_date(end)}  •  {days} day(s)"
            ),
            headers=headers,
            rows=rows,
            summary=[
                ("Period", f"{start.isoformat()} to {end.isoformat()}"),
                ("Records", str(len(records))),
                ("Total hours", clock.format_duration(totals_minutes)),
            ],
            totals=["TOTAL", "", "", "", "", "", clock.format_duration(totals_minutes), f"{len(rows)} row(s)"],
        )

    # -- employee sheet ------------------------------------------------------
    def employee_report(self, employee: Employee, window: WorkWeek | None = None) -> ReportResult:
        clock = self.clock
        window = window or clock.week_bounds()
        records = self.attendance.employee_history(employee, window.start, window.end)
        progress = self.attendance.week_progress(employee, window)
        headers = ["Date", "Time In", "Time Out", "Hours", "Status", "Source"]
        rows: list[list] = []
        totals_minutes = 0
        for record in records:
            minutes = record.duration_minutes or 0
            if record.is_open:
                minutes = TimeUtils.elapsed_minutes(
                    TimeUtils.parse(record.time_in) or clock.now(), clock.now()
                )
            totals_minutes += minutes
            rows.append(
                [
                    clock.format_date(record.work_date),
                    clock.format_time(record.time_in),
                    clock.format_time(record.time_out),
                    clock.format_duration(minutes),
                    record.status_label,
                    "Manual" if record.note.startswith("Added") else "QR",
                ]
            )
        return ReportResult(
            title=f"Attendance Report - {employee.full_name}",
            subtitle=(
                f"{employee.employee_code}  •  "
                f"{clock.format_date(window.start)} to {clock.format_date(window.end)}"
            ),
            headers=headers,
            rows=rows,
            summary=[
                ("Employee", employee.full_name),
                ("Employee ID", employee.employee_code),
                ("Department", employee.department or "-"),
                ("Position", employee.position or "-"),
                ("Hours in period", clock.format_duration(totals_minutes)),
                ("Weekly goal", progress.goal_text),
                ("Progress", f"{progress.percent}%"),
                ("Goal reached", "Yes" if progress.goal_reached else "No"),
            ],
            totals=["TOTAL", "", "", clock.format_duration(totals_minutes), "", ""],
        )

    # -- audit ---------------------------------------------------------------
    def audit_logs(self, limit: int = 500, action: str = "", search: str = "") -> ReportResult:
        logs: list[AuditLog] = self.repos.audit.list_recent(limit=limit, action=action, search=search)
        headers = ["When (UTC)", "Administrator", "Action", "Entity", "Description", "Old Value", "New Value", "Reason"]
        rows = [
            [
                log.created_at.replace("T", " "),
                log.admin_username,
                log.action_label,
                f"{log.entity_type}:{log.entity_id}" if log.entity_type else "-",
                log.description,
                log.old_value or "-",
                log.new_value or "-",
                log.reason or "-",
            ]
            for log in logs
        ]
        return ReportResult(
            title="Audit Log",
            subtitle=f"{self.settings.settings.organization_name}  •  {len(logs)} entry(ies)",
            headers=headers,
            rows=rows,
            summary=[("Entries", str(len(logs)))],
        )

    # -- exports -------------------------------------------------------------
    def write_csv(self, report: ReportResult, destination: Path | str) -> Path:
        path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle)
            writer.writerow([report.title])
            writer.writerow([report.subtitle])
            writer.writerow([])
            for label, value in report.summary:
                writer.writerow([label, value])
            writer.writerow([])
            writer.writerow(report.headers)
            for row in report.rows:
                writer.writerow(row)
            if report.totals:
                writer.writerow([])
                writer.writerow(report.totals)
        return path

    def csv_bytes(self, report: ReportResult) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow([report.title])
        writer.writerow([report.subtitle])
        writer.writerow([])
        for label, value in report.summary:
            writer.writerow([label, value])
        writer.writerow([])
        writer.writerow(report.headers)
        for row in report.rows:
            writer.writerow(row)
        if report.totals:
            writer.writerow([])
            writer.writerow(report.totals)
        return buffer.getvalue().encode("utf-8-sig")

    def write_xlsx(self, report: ReportResult, destination: Path | str) -> Path:
        if not XLSX_AVAILABLE:
            raise RuntimeError(
                "XLSX export needs the 'openpyxl' package. Use CSV export instead."
            )
        path = Path(destination)
        path.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Report"

        header_fill = PatternFill("solid", fgColor="1E293B")
        title_font = Font(bold=True, size=14, color="0F172A")
        header_font = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin", color="CBD5E1")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        sheet.cell(row=1, column=1, value=report.title).font = title_font
        sheet.cell(row=2, column=1, value=report.subtitle)
        row_index = 4
        for label, value in report.summary:
            sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
            sheet.cell(row=row_index, column=2, value=value)
            row_index += 1
        row_index += 1

        header_row = row_index
        for column, title in enumerate(report.headers, start=1):
            cell = sheet.cell(row=header_row, column=column, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for offset, row in enumerate(report.rows, start=1):
            for column, value in enumerate(row, start=1):
                cell = sheet.cell(row=header_row + offset, column=column, value=value)
                cell.border = border

        if report.totals:
            total_row = header_row + len(report.rows) + 2
            for column, value in enumerate(report.totals, start=1):
                cell = sheet.cell(row=total_row, column=column, value=value)
                cell.font = Font(bold=True)
                cell.border = border

        widths = [
            max(12, min(38, max([len(str(report.headers[i]))] + [
                len(str(row[i])) for row in report.rows if i < len(row)
            ]) + 2))
            for i in range(len(report.headers))
        ]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

        workbook.save(str(path))
        return path

    # -- filenames -----------------------------------------------------------
    def suggest_filename(self, report: ReportResult, extension: str) -> str:
        clock = self.clock
        slug = "".join(ch if ch.isalnum() else "-" for ch in report.title.lower())
        slug = "-".join(part for part in slug.split("-") if part)[:60]
        return f"{slug}_{clock.export_stamp()}.{extension}"


def clock_format(minutes: int) -> str:
    return TimeUtils.format_duration(minutes)


__all__ = ["ReportResult", "ReportService", "ReportRow", "XLSX_AVAILABLE"]