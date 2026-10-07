"""Report templates: a reusable, user-editable column layout.

A template decides which columns an export contains, in what order, what each
one is called, how wide it is, and whether a duration is written as text
(``8h 12m``) or as a decimal number (``8.2``) so Excel or Sheets can total it.

The same template drives both directions:

* :func:`apply_template` reshapes a :class:`~app.services.report_service.ReportResult`
  for export (CSV, XLSX, Google Sheets).
* :func:`parse_tabular` reads a CSV or XLSX file back using the same column
  headings, so a file Vardia exported can be edited elsewhere and imported.

Templates are stored as JSON in the settings table under
``report_templates``, keyed by name. Built-in templates are provided and can
be duplicated and edited; nothing is ever silently discarded.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

#: Format for a cell that holds a duration.
FORMAT_TEXT = "text"      # "8h 12m" - what a person reads
FORMAT_HOURS = "hours"    # 8.2 - what a spreadsheet can sum

#: Report kinds a template can be bound to. Mirrors the Reports page.
KIND_DAILY = "daily"
KIND_WEEKLY = "weekly"
KIND_MONTHLY = "monthly"
KIND_RANGE = "range"
KIND_EMPLOYEE = "employee"
KIND_AUDIT = "audit"
REPORT_KINDS: tuple[str, ...] = (
    KIND_DAILY,
    KIND_WEEKLY,
    KIND_MONTHLY,
    KIND_RANGE,
    KIND_EMPLOYEE,
    KIND_AUDIT,
)

TEMPLATES_SETTING_KEY = "report_templates"


# ---------------------------------------------------------------------------
# duration helpers
# ---------------------------------------------------------------------------
_DURATION_RE = re.compile(
    r"^\s*(?:(?P<h>\d+)\s*h)?\s*(?:(?P<m>\d+)\s*m)?\s*$", re.IGNORECASE
)


def duration_to_hours(value: Any) -> Any:
    """Convert a duration to decimal hours.

    Accepts ``"8h 12m"``, ``"8.23"``, ``8.23`` and ``8`` so a file exported
    with numeric hours can be read back in. Returns ``None`` when the value is
    not a duration, so callers can tell "no data" from "zero".
    """
    if value is None:
        return None
    if isinstance(value, bool):  # bool is an int; never a duration
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text == "-":
        return None
    match = _DURATION_RE.match(text)
    if match:
        hours = match.group("h")
        minutes = match.group("m")
        if hours is not None or minutes is not None:
            return int(hours or 0) + (int(minutes or 0) / 60.0)
        return None
    # A plain decimal, as written by a spreadsheet or by our own export.
    try:
        return float(text)
    except ValueError:
        return None


def hours_to_duration(value: Any) -> str:
    """Convert ``8.2`` (or ``"8.2"``) to ``"8h 12m"``. Empty input -> ``""``."""
    if value is None:
        return ""
    if isinstance(value, str) and not value.strip():
        return ""
    try:
        total = float(value)
    except (TypeError, ValueError):
        return str(value)
    if total < 0:
        total = 0.0
    minutes = int(round(total * 60))
    return f"{minutes // 60}h {minutes % 60:02d}m"


# ---------------------------------------------------------------------------
# model
# ---------------------------------------------------------------------------
@dataclass
class TemplateColumn:
    """One column in a template.

    ``key`` is the report's own heading (the identity used when matching).
    ``heading`` is what the exported file shows, so an admin can rename
    "Employee ID" to "Staff No." without breaking the mapping.
    """

    key: str
    heading: str = ""
    width: int = 0          # 0 = size it from the content
    number_format: str = FORMAT_TEXT
    included: bool = True

    def __post_init__(self) -> None:
        if not self.heading:
            self.heading = self.key
        if self.number_format not in (FORMAT_TEXT, FORMAT_HOURS):
            self.number_format = FORMAT_TEXT
        self.width = max(0, min(60, int(self.width or 0)))

    def to_dict(self) -> dict[str, Any]:
        return {
            "key": self.key,
            "heading": self.heading,
            "width": self.width,
            "number_format": self.number_format,
            "included": self.included,
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "TemplateColumn":
        return cls(
            key=str(data.get("key", "")),
            heading=str(data.get("heading", "") or ""),
            width=int(data.get("width", 0) or 0),
            number_format=str(data.get("number_format", FORMAT_TEXT)),
            included=bool(data.get("included", True)),
        )


@dataclass
class ReportTemplate:
    """A named column layout bound to a report kind."""

    name: str
    kind: str = KIND_DAILY
    columns: list[TemplateColumn] = field(default_factory=list)
    include_summary: bool = True
    include_totals: bool = True
    builtin: bool = False

    def active_columns(self) -> list[TemplateColumn]:
        return [column for column in self.columns if column.included]

    def __post_init__(self) -> None:
        self.kind = self.kind if self.kind in REPORT_KINDS else KIND_DAILY
        self.columns = [c for c in self.columns if c.key]

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "kind": self.kind,
            "columns": [column.to_dict() for column in self.columns],
            "include_summary": self.include_summary,
            "include_totals": self.include_totals,
            "builtin": self.builtin,
        }

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> "ReportTemplate":
        raw_columns = data.get("columns") or []
        return cls(
            name=str(data.get("name", "")).strip() or "Untitled",
            kind=str(data.get("kind", KIND_DAILY)),
            columns=[
                TemplateColumn.from_dict(item)
                for item in raw_columns
                if isinstance(item, dict)
            ],
            include_summary=bool(data.get("include_summary", True)),
            include_totals=bool(data.get("include_totals", True)),
            builtin=bool(data.get("builtin", False)),
        )

    def duplicate(self, new_name: str) -> "ReportTemplate":
        clone = ReportTemplate.from_dict(self.to_dict())
        clone.name = new_name
        clone.builtin = False
        return clone

    def find(self, key: str) -> TemplateColumn | None:
        """Match a column by its key or its heading, ignoring case and spacing."""
        lowered = (key or "").strip().lower()
        for column in self.columns:
            if column.key.strip().lower() == lowered:
                return column
        for column in self.columns:
            if column.heading.strip().lower() == lowered:
                return column
        return None


# ---------------------------------------------------------------------------
# built-ins
# ---------------------------------------------------------------------------
def _cols(*specs: tuple[str, str, str, int]) -> list[TemplateColumn]:
    return [
        TemplateColumn(key=key, heading=heading, number_format=fmt, width=width)
        for key, heading, fmt, width in specs
    ]


def builtin_templates() -> list[ReportTemplate]:
    """The layouts shipped with VardiaShift, one per report kind."""
    return [
        ReportTemplate(
            name="Daily attendance (standard)",
            kind=KIND_DAILY,
            columns=_cols(
                ("Employee ID", "Employee ID", FORMAT_TEXT, 14),
                ("Full Name", "Full Name", FORMAT_TEXT, 22),
                ("Department", "Department", FORMAT_TEXT, 16),
                ("Time In", "Time In", FORMAT_TEXT, 12),
                ("Time Out", "Time Out", FORMAT_TEXT, 12),
                ("Hours", "Hours", FORMAT_TEXT, 10),
                ("Status", "Status", FORMAT_TEXT, 14),
                ("Source", "Source", FORMAT_TEXT, 14),
            ),
            builtin=True,
        ),
        ReportTemplate(
            name="Daily attendance (hours as numbers)",
            kind=KIND_DAILY,
            columns=_cols(
                ("Employee ID", "Employee ID", FORMAT_TEXT, 14),
                ("Full Name", "Full Name", FORMAT_TEXT, 22),
                ("Department", "Department", FORMAT_TEXT, 16),
                ("Time In", "Time In", FORMAT_TEXT, 12),
                ("Time Out", "Time Out", FORMAT_TEXT, 12),
                ("Hours", "Hours Worked", FORMAT_HOURS, 14),
                ("Status", "Status", FORMAT_TEXT, 14),
                ("Source", "Source", FORMAT_TEXT, 14),
            ),
            builtin=True,
        ),
        ReportTemplate(
            name="Payroll sheet",
            kind=KIND_RANGE,
            columns=_cols(
                ("Date", "Date", FORMAT_TEXT, 12),
                ("Employee ID", "Employee ID", FORMAT_TEXT, 14),
                ("Full Name", "Employee Name", FORMAT_TEXT, 24),
                ("Department", "Department", FORMAT_TEXT, 16),
                ("Time In", "Clock In", FORMAT_TEXT, 12),
                ("Time Out", "Clock Out", FORMAT_TEXT, 12),
                ("Hours", "Hours", FORMAT_HOURS, 12),
                ("Status", "Status", FORMAT_TEXT, 14),
            ),
            builtin=True,
        ),
        ReportTemplate(
            name="Weekly summary",
            kind=KIND_WEEKLY,
            columns=_cols(
                ("Employee ID", "Employee ID", FORMAT_TEXT, 14),
                ("Full Name", "Full Name", FORMAT_TEXT, 22),
                ("Department", "Department", FORMAT_TEXT, 16),
                ("Hours Worked", "Hours Worked", FORMAT_HOURS, 14),
                ("Weekly Goal", "Weekly Goal", FORMAT_HOURS, 14),
                ("Remaining", "Remaining", FORMAT_HOURS, 14),
                ("Overtime", "Overtime", FORMAT_HOURS, 12),
                ("Days", "Days", FORMAT_TEXT, 10),
                ("Status", "Status", FORMAT_TEXT, 14),
            ),
            builtin=True,
        ),
        ReportTemplate(
            name="Compact (id and hours only)",
            kind=KIND_DAILY,
            columns=_cols(
                ("Employee ID", "Employee ID", FORMAT_TEXT, 14),
                ("Full Name", "Full Name", FORMAT_TEXT, 24),
                ("Hours", "Hours", FORMAT_HOURS, 12),
            ),
            include_summary=False,
            include_totals=False,
            builtin=True,
        ),
        ReportTemplate(
            name="Audit log",
            kind=KIND_AUDIT,
            columns=_cols(
                ("When (UTC)", "When", FORMAT_TEXT, 22),
                ("Administrator", "Administrator", FORMAT_TEXT, 18),
                ("Action", "Action", FORMAT_TEXT, 22),
                ("Entity", "Entity", FORMAT_TEXT, 18),
                ("Description", "Description", FORMAT_TEXT, 40),
                ("Old Value", "Old Value", FORMAT_TEXT, 30),
                ("New Value", "New Value", FORMAT_TEXT, 30),
                ("Reason", "Reason", FORMAT_TEXT, 30),
            ),
            builtin=True,
        ),
    ]


# ---------------------------------------------------------------------------
# storage
# ---------------------------------------------------------------------------
def load_templates(settings) -> list[ReportTemplate]:
    """All templates: the built-ins plus any saved by the administrator."""
    templates: list[ReportTemplate] = []
    raw = ""
    try:
        raw = settings.get(TEMPLATES_SETTING_KEY, "") or ""
    except Exception:  # pragma: no cover - settings unavailable
        raw = ""
    if raw:
        try:
            payload = json.loads(raw)
            if isinstance(payload, list):
                for item in payload:
                    if isinstance(item, dict):
                        templates.append(ReportTemplate.from_dict(item))
        except (ValueError, TypeError):
            # A corrupt value must not stop the app starting; the built-ins
            # still work, so fall through to them.
            templates = []

    names = {template.name for template in templates}
    for builtin in builtin_templates():
        if builtin.name not in names:
            templates.append(builtin)
    return templates


def save_templates(settings, templates: Iterable[ReportTemplate], admin_username: str = "system") -> None:
    """Persist the administrator's own templates (built-ins are not stored)."""
    custom = [template.to_dict() for template in templates if not template.builtin]
    settings.set(TEMPLATES_SETTING_KEY, json.dumps(custom), admin_username)


def templates_for_kind(settings, kind: str) -> list[ReportTemplate]:
    return [template for template in load_templates(settings) if template.kind == kind]


def template_by_name(settings, name: str) -> ReportTemplate | None:
    for template in load_templates(settings):
        if template.name == name:
            return template
    return None


def resolve_template(settings, name: str | None, kind: str, headers: Iterable[str]):
    """Find the template to use, falling back to matching the report itself.

    Returns ``(template, warnings)``. A template that does not cover every
    report column still applies: the uncovered columns are reported so the
    user is told what will be dropped rather than losing data silently.
    """
    headers = list(headers)
    template = template_by_name(settings, name) if name else None
    warnings: list[str] = []

    if template is None:
        # No template chosen: mirror the report as-is so default behaviour
        # is unchanged.
        template = ReportTemplate(
            name="(report default)",
            kind=kind,
            columns=[TemplateColumn(key=header) for header in headers],
            builtin=True,
        )
        return template, warnings

    if template.kind != kind:
        warnings.append(
            f"Template '{template.name}' is built for {template.kind} reports, "
            f"not {kind}."
        )

    covered = {column.key.strip().lower() for column in template.active_columns()}
    for header in headers:
        if header.strip().lower() not in covered:
            warnings.append(
                f"'{header}' is not part of this template and will be left out."
            )
    for column in template.active_columns():
        if not any(column.key.strip().lower() == h.strip().lower() for h in headers):
            warnings.append(
                f"'{column.heading}' is in the template but not in this report, "
                "so the column stays empty."
            )
    return template, warnings


# ---------------------------------------------------------------------------
# export
# ---------------------------------------------------------------------------
@dataclass
class TemplateReport:
    """A report reshaped by a template, ready for CSV/XLSX/Sheets."""

    title: str
    subtitle: str
    headers: list[str]
    rows: list[list]
    summary: list[tuple[str, str]] = field(default_factory=list)
    totals: list[str] = field(default_factory=list)
    widths: list[int] = field(default_factory=list)

    def to_report_result(self):
        """Wrap back into a ReportResult so existing exporters work unchanged."""
        from app.services.report_service import ReportResult

        return ReportResult(
            title=self.title,
            subtitle=self.subtitle,
            headers=self.headers,
            rows=self.rows,
            summary=self.summary,
            totals=self.totals,
        )


def apply_template(report, template: ReportTemplate) -> TemplateReport:
    """Reshape ``report`` into the column order and formats of ``template``."""
    source_headers = [str(h) for h in report.headers]
    lookup = {header.strip().lower(): index for index, header in enumerate(source_headers)}

    headers: list[str] = []
    widths: list[int] = []
    rows: list[list] = []

    for column in template.active_columns():
        headers.append(column.heading)
        widths.append(column.width)
        index = lookup.get(column.key.strip().lower())
        new_row: list[Any] = []
        for row in report.rows:
            if index is None:
                new_row.append(None)
                continue
            raw = row[index] if index < len(row) else None
            if column.number_format == FORMAT_HOURS:
                hours = duration_to_hours(raw)
                # A duration becomes a number. A placeholder ("-") or an
                # empty cell becomes None so it is never mistaken for 0
                # worked. Anything else that is not a duration is kept as
                # text, so a mislabelled column cannot silently lose data.
                if hours is not None:
                    new_row.append(hours)
                elif raw in (None, "") or _cell_to_text(raw) == "-":
                    new_row.append(None)
                else:
                    new_row.append(raw)
            else:
                new_row.append(raw)
        rows.append(new_row)

    # Transpose: we built column-major, reports are row-major.
    transposed = [
        [rows[column][row] for column in range(len(rows))]
        for row in range(len(report.rows))
    ]

    totals: list[str] = []
    if template.include_totals and report.totals:
        totals = []
        for column in template.active_columns():
            source_index = lookup.get(column.key.strip().lower())
            totals.append(
                _reshape_totals_cell(
                    column,
                    report.totals[source_index] if source_index is not None else None,
                )
            )

    return TemplateReport(
        title=report.title,
        subtitle=report.subtitle,
        headers=headers,
        rows=transposed,
        summary=list(report.summary) if template.include_summary else [],
        totals=totals,
        widths=widths,
    )


def _reshape_totals_cell(column: TemplateColumn, value: Any) -> str:
    """Format one totals cell for ``column``."""
    if column.number_format == FORMAT_HOURS:
        hours = duration_to_hours(value)
        return "" if hours is None else f"{hours:.2f}"
    return "" if value is None else str(value)


def to_csv_bytes(shaped: TemplateReport) -> bytes:
    """CSV bytes for a templated report, formula-injection safe."""
    from app.services.report_service import sanitize_spreadsheet_cell

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([shaped.title])
    writer.writerow([shaped.subtitle])
    writer.writerow([])
    for label, value in shaped.summary:
        writer.writerow([label, sanitize_spreadsheet_cell(value)])
    writer.writerow([])
    writer.writerow([sanitize_spreadsheet_cell(cell) for cell in shaped.headers])
    for row in shaped.rows:
        writer.writerow([sanitize_spreadsheet_cell(cell) for cell in row])
    if shaped.totals:
        writer.writerow([])
        writer.writerow([sanitize_spreadsheet_cell(cell) for cell in shaped.totals])
    return buffer.getvalue().encode("utf-8-sig")


def to_sheet_values(shaped: TemplateReport) -> list[list]:
    """Rows for the Google Sheets API, matching :func:`to_csv_bytes`."""
    from app.services.report_service import sanitize_spreadsheet_cell

    def cell(value) -> str:
        clean = sanitize_spreadsheet_cell(value)
        if clean is None:
            return ""
        if isinstance(clean, float):
            return f"{clean:.2f}"
        return str(clean)

    rows: list[list] = [[shaped.title], [shaped.subtitle], []]
    for label, value in shaped.summary:
        rows.append([cell(label), cell(value)])
    rows.append([])
    rows.append([cell(header) for header in shaped.headers])
    for row in shaped.rows:
        rows.append([cell(value) for value in row])
    if shaped.totals:
        rows.append([])
        rows.append([cell(value) for value in shaped.totals])
    return rows


def write_xlsx(shaped: TemplateReport, destination: Path | str) -> Path:
    """Write a templated report to XLSX with the template's column widths."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - openpyxl ships with VardiaShift
        raise RuntimeError(
            "XLSX export needs the 'openpyxl' package. Use CSV export instead."
        ) from exc

    from app.services.report_service import sanitize_spreadsheet_cell

    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    header_fill = PatternFill("solid", fgColor="1E293B")
    title_font = Font(bold=True, size=14, color="0F172A")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.cell(row=1, column=1, value=shaped.title).font = title_font
    sheet.cell(row=2, column=1, value=shaped.subtitle)

    row_index = 4
    for label, value in shaped.summary:
        sheet.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=row_index, column=2, value=sanitize_spreadsheet_cell(value))
        row_index += 1

    header_row = row_index
    for column, heading in enumerate(shaped.headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=heading)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for offset, row in enumerate(shaped.rows, start=1):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(
                row=header_row + offset,
                column=column,
                value=sanitize_spreadsheet_cell(value),
            )
            cell.border = border

    if shaped.totals:
        total_row = header_row + len(shaped.rows) + 2
        for column, value in enumerate(shaped.totals, start=1):
            cell = sheet.cell(
                row=total_row, column=column, value=sanitize_spreadsheet_cell(value)
            )
            cell.font = Font(bold=True)
            cell.border = border

    for index in range(len(shaped.headers)):
        letter = get_column_letter(index + 1)
        if shaped.widths and shaped.widths[index]:
            sheet.column_dimensions[letter].width = shaped.widths[index]
        else:
            values = [str(shaped.headers[index])] + [
                str(row[index]) for row in shaped.rows if index < len(row) and row[index] is not None
            ]
            sheet.column_dimensions[letter].width = max(
                12, min(38, max(len(value) for value in values) + 2)
            )

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    workbook.save(str(path))
    return path


# ---------------------------------------------------------------------------
# import
# ---------------------------------------------------------------------------
@dataclass
class ImportRow:
    """One parsed data row, with the values keyed by template column."""

    line: int
    values: dict[str, Any]

    def get(self, key: str, default: Any = "") -> Any:
        return self.values.get(key, default)


@dataclass
class ImportPreview:
    """Result of reading a file, ready to show before anything is written."""

    rows: list[ImportRow] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    headers_seen: list[str] = field(default_factory=list)
    sheet_name: str = ""

    @property
    def ok(self) -> bool:
        return bool(self.rows) and not self.errors


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_header_row(matrix: list[list[Any]], template: ReportTemplate) -> int:
    """Locate the header row: the line that matches the most template keys."""
    wanted = {column.key.strip().lower() for column in template.active_columns()}
    if not wanted:
        return -1
    best_index, best_score = -1, 0
    for index, row in enumerate(matrix[:30]):
        present = {_cell_to_text(value).strip().lower() for value in row}
        score = len(wanted & present)
        if score > best_score:
            best_index, best_score = index, score
    return best_index if best_score else -1


def parse_tabular(source: Path | str | bytes, template: ReportTemplate) -> ImportPreview:
    """Read a CSV or XLSX file back using ``template``'s column headings.

    Header matching is deliberately forgiving: a cell matches a template
    column when it equals the column's key *or* its heading, ignoring case and
    surrounding spaces. Unmatched template columns come back empty and
    unmatched file columns are reported rather than dropped silently.
    """
    matrix, sheet_name = _read_matrix(source)
    preview = ImportPreview(sheet_name=sheet_name)

    if not matrix:
        preview.errors.append("The file is empty.")
        return preview

    header_index = _find_header_row(matrix, template)
    if header_index < 0:
        preview.errors.append(
            "Could not find a header row. Expected a row containing at least "
            "one of: "
            + ", ".join(sorted({column.heading for column in template.active_columns()}))
        )
        return preview

    header_row = matrix[header_index]
    preview.headers_seen = [_cell_to_text(value) for value in header_row if _cell_to_text(value)]

    # Map file column index -> template column
    mapping: dict[int, TemplateColumn] = {}
    matched_keys: set[str] = set()
    for index, value in enumerate(header_row):
        text = _cell_to_text(value).strip().lower()
        if not text:
            continue
        column = template.find(text)
        if column is None:
            for candidate in template.active_columns():
                if candidate.heading.strip().lower() == text and candidate.key not in matched_keys:
                    column = candidate
                    break
        if column is not None and column.key not in matched_keys:
            mapping[index] = column
            matched_keys.add(column.key)

    if not mapping:
        preview.errors.append("No column in the file matched this template.")
        return preview

    for column in template.active_columns():
        if column.key not in matched_keys:
            preview.warnings.append(
                f"Column '{column.heading}' is not in the file; its values stay empty."
            )

    for index, value in enumerate(header_row):
        text = _cell_to_text(value)
        if text and index not in mapping:
            preview.warnings.append(f"Ignoring unknown column '{text}'.")

    # An export ends with a blank row then a TOTAL row. Importing that TOTAL
    # row as data would create a bogus attendance record, so it is skipped.
    totals_markers = {"total", "totals", "grand total", "sum"}

    for offset, row in enumerate(matrix[header_index + 1 :], start=header_index + 2):
        cells = [_cell_to_text(value) for value in row]
        if not any(cells):
            continue
        first = next((text for text in cells if text), "")
        if first.strip().lower() in totals_markers:
            continue
        values: dict[str, Any] = {}
        for index, column in mapping.items():
            raw = row[index] if index < len(row) else None
            if column.number_format == FORMAT_HOURS:
                values[column.key] = duration_to_hours(raw)
            else:
                values[column.key] = _cell_to_text(raw)
        preview.rows.append(ImportRow(line=offset, values=values))

    if not preview.rows:
        preview.errors.append("The file has a header row but no data rows.")
    return preview


def _read_matrix(source: Path | str | bytes) -> tuple[list[list[Any]], str]:
    """Return ``(rows, sheet_name)`` from CSV bytes, an XLSX path, or CSV text."""
    if isinstance(source, bytes):
        text = source.decode("utf-8-sig", errors="replace")
        return _parse_csv(text), ""

    path = Path(source)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - ships with VardiaShift
            raise RuntimeError(
                "Reading XLSX needs the 'openpyxl' package."
            ) from exc
        workbook = load_workbook(str(path), data_only=True, read_only=True)
        sheet = workbook.active
        rows = [
            [cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        workbook.close()
        return rows, sheet.title

    # Unknown extension: try CSV, which is the friendlier failure.
    try:
        return _parse_csv(path.read_text(encoding="utf-8-sig", errors="replace")), ""
    except OSError as exc:
        raise RuntimeError(f"Could not read {path.name}: {exc}") from exc


def _parse_csv(text: str) -> list[list[str]]:
    return [list(row) for row in csv.reader(io.StringIO(text))]


__all__ = [
    "FORMAT_HOURS",
    "FORMAT_TEXT",
    "ImportPreview",
    "ImportRow",
    "KIND_AUDIT",
    "KIND_DAILY",
    "KIND_EMPLOYEE",
    "KIND_MONTHLY",
    "KIND_RANGE",
    "KIND_WEEKLY",
    "REPORT_KINDS",
    "ReportTemplate",
    "TEMPLATES_SETTING_KEY",
    "TemplateColumn",
    "TemplateReport",
    "apply_template",
    "builtin_templates",
    "duration_to_hours",
    "hours_to_duration",
    "load_templates",
    "parse_tabular",
    "resolve_template",
    "save_templates",
    "template_by_name",
    "templates_for_kind",
    "to_csv_bytes",
    "to_sheet_values",
    "write_xlsx",
]