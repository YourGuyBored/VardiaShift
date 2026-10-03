"""Report generation and CSV / XLSX export."""

from __future__ import annotations

import csv
from datetime import datetime, time, timedelta
from pathlib import Path

import pytest

from app.services.report_service import XLSX_AVAILABLE


def work(context, employee, start_h, out_h, offset=0, frozen=None):
    clock = frozen
    context.attendance.time_in(employee, clock.at(start_h, 0, offset))
    context.attendance.time_out(employee, clock.at(out_h, 0, offset))


@pytest.fixture
def team(context, admin, frozen):
    juan = context.employees.create("EMP-001", "Juan Dela Cruz", "IT", "Student Assistant", admin_username="admin")
    maria = context.employees.create("EMP-002", "Maria Santos", "HR", "Coordinator", admin_username="admin")
    pedro = context.employees.create("EMP-003", "Pedro Reyes", "IT", "Assistant", admin_username="admin")

    work(context, juan, 9, 17, 0, frozen)          # 8h
    work(context, juan, 9, 17, 1, frozen)          # 8h  -> 16h
    work(context, maria, 8, 16, 0, frozen)         # 8h
    work(context, pedro, 8, 15, 0, frozen)         # 7h
    return juan, maria, pedro


# -- daily -------------------------------------------------------------------
def test_daily_report_lists_all_active_employees(context, admin, team, frozen):
    report = context.reports.daily_report(frozen.reference.date(), employee_id=None)
    codes = {row[0] for row in report.rows}
    assert codes == {"EMP-001", "EMP-002", "EMP-003"}
    assert "Daily Attendance Report" in report.title
    assert report.headers[0] == "Employee ID"


def test_daily_report_includes_not_in_rows(context, admin, team, frozen):
    friday = frozen.at(23, 0, 4).date()   # nobody worked on day 4 of the week
    report = context.reports.daily_report(friday, employee_id=None)
    assert report.rows
    assert {row[-1] for row in report.rows} == {"Not In"}
    assert {row[0] for row in report.rows} == {"EMP-001", "EMP-002", "EMP-003"}


def test_daily_report_computes_hours(context, admin, frozen):
    emp = context.employees.create("EMP-001", "Juan Dela Cruz", admin_username="admin")
    work(context, emp, 9, 17, 0, frozen)
    report = context.reports.daily_report(frozen.reference.date(), emp.employee_id)
    assert report.rows[0][5] == "8h 00m"
    assert report.rows[0][6] == "Timed Out"


def test_daily_report_filtered_by_employee(context, admin, team, frozen):
    report = context.reports.daily_report(frozen.reference.date(), employee_id=team[0].employee_id)
    assert len(report.rows) == 1
    assert report.rows[0][1] == "Juan Dela Cruz"


def test_daily_report_summary(context, admin, team, frozen):
    report = context.reports.daily_report(frozen.reference.date())
    labels = dict(report.summary)
    assert labels["Employees present"] == "3"
    assert labels["Timed out"] == "3"
    assert labels["Total hours"] == "23h 00m"  # Monday only: 8h + 8h + 7h


def test_daily_report_handles_open_session(context, admin, frozen):
    emp = context.employees.create("EMP-001", "Live Session", admin_username="admin")
    context.attendance.time_in(emp, frozen.at(8, 0))
    report = context.reports.daily_report(frozen.reference.date(), emp.employee_id)
    assert report.rows[0][3] == "08:00 AM"
    assert report.rows[0][4] == "-"
    assert report.rows[0][6] == "Working"


# -- weekly ------------------------------------------------------------------
def test_weekly_report_columns(context, admin, team, frozen):
    report = context.reports.weekly_report(context.clock.week_bounds())
    assert report.headers == [
        "Employee ID", "Full Name", "Department", "Hours Worked", "Weekly Goal",
        "Remaining", "Overtime", "Days", "Status",
    ]
    assert "Weekly" in report.title or "W" in report.title


def test_weekly_report_values(context, admin, team, frozen):
    report = context.reports.weekly_report(context.clock.week_bounds())
    rows = {row[0]: row for row in report.rows}
    assert rows["EMP-001"][3] == "16h 00m"
    assert rows["EMP-001"][4] == "30h 00m"
    assert rows["EMP-001"][5] == "14h 00m"
    assert rows["EMP-001"][6] == "-"
    assert rows["EMP-001"][7] == "2"
    assert rows["EMP-001"][8] == "In progress"
    assert rows["EMP-002"][3] == "8h 00m"


def test_weekly_report_overtime_column(context, admin, frozen):
    emp = context.employees.create("EMP-001", "Overtime Person", admin_username="admin")
    context.employees.set_weekly_goal(emp, 6, "admin")
    work(context, emp, 9, 17, 0, frozen)  # 8h vs 6h goal
    report = context.reports.weekly_report(context.clock.week_bounds())
    row = report.rows[0]
    assert row[6] == "+2h 00m"
    assert "Goal reached" in row[8]


def test_weekly_report_goal_reached(context, admin, frozen):
    emp = context.employees.create("EMP-001", "Goal Person", admin_username="admin")
    context.employees.set_weekly_goal(emp, 8, "admin")
    work(context, emp, 9, 17, 0, frozen)
    report = context.reports.weekly_report(context.clock.week_bounds())
    assert report.rows[0][8] == "Goal reached"


def test_weekly_report_totals(context, admin, team, frozen):
    report = context.reports.weekly_report(context.clock.week_bounds())
    assert report.totals[3] == "31h 00m"
    assert report.totals[4] == "90h 00m"


def test_weekly_report_uses_configured_goal(context, admin, team, frozen):
    context.settings.set("default_weekly_goal_hours", 20, "admin")
    report = context.reports.weekly_report(context.clock.week_bounds())
    rows = {row[0]: row for row in report.rows}
    assert rows["EMP-001"][4] == "20h 00m"


def test_weekly_report_last_week(context, admin, team, frozen):
    report = context.reports.weekly_report(context.clock.last_week_bounds())
    assert report.is_empty or all(row[3] == "0h 00m" for row in report.rows)


# -- monthly -----------------------------------------------------------------
def test_monthly_report_details(context, admin, team, frozen):
    from app.utils.time_utils import TimeUtils

    window = TimeUtils().month_bounds(frozen.reference.date())
    report = context.reports.monthly_report(window)
    assert report.headers[0] == "Date"
    assert report.rows
    assert report.totals[5] == "31h 00m"
    assert any("Juan Dela Cruz" in line for line, _ in report.summary)


def test_monthly_report_is_empty_for_future_month(context, admin, team, frozen):
    from datetime import timedelta

    start = frozen.reference.date() + timedelta(days=10)
    report = context.reports.range_report(start, start + timedelta(days=5))
    assert report.is_empty


# -- custom range ------------------------------------------------------------
def test_range_report(context, admin, team, frozen):
    start = frozen.at(0, 0).date() - timedelta(days=1)
    end = frozen.at(0, 0).date() + timedelta(days=1)
    report = context.reports.range_report(start, end)
    assert report.headers[0] == "Date"
    assert len(report.rows) >= 3


def test_employee_report(context, admin, team, frozen):
    juan = team[0]
    report = context.reports.employee_report(juan, context.clock.week_bounds())
    assert report.title.endswith(juan.full_name)
    summary = dict(report.summary)
    assert summary["Employee ID"] == "EMP-001"
    assert summary["Department"] == "IT"
    assert summary["Hours in period"] == "16h 00m"
    assert summary["Progress"] == "53.3%"


# -- audit -------------------------------------------------------------------
def test_audit_report(context, admin, team, frozen):
    context.employees.deactivate(team[2], "admin", reason="Left")
    report = context.reports.audit_logs()
    assert report.title == "Audit Log"
    assert any("Deactivated" in str(row[4]) for row in report.rows)


def test_audit_report_filtered(context, admin, team, frozen):
    context.employees.deactivate(team[2], "admin")
    report = context.reports.audit_logs(action="employee_deactivate")
    assert report.rows
    assert all(row[2] == "Employee deactivated" for row in report.rows)


# -- exports -----------------------------------------------------------------
def test_csv_export(context, admin, team, frozen, tmp_path):
    report = context.reports.weekly_report(context.clock.week_bounds())
    path = context.reports.write_csv(report, tmp_path / "weekly.csv")
    assert path.is_file()
    text = path.read_text(encoding="utf-8-sig")
    assert "Weekly Attendance Report" in text
    assert "EMP-001" in text
    assert "16h 00m" in text


def test_csv_is_parseable(context, admin, team, frozen, tmp_path):
    report = context.reports.daily_report(frozen.reference.date())
    path = context.reports.write_csv(report, tmp_path / "daily.csv")
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert any("Employee ID" in cell for row in rows for cell in row)


def test_csv_bytes(context, admin, team, frozen):
    from app.utils.time_utils import TimeUtils

    window = TimeUtils().month_bounds(frozen.reference.date())
    report = context.reports.monthly_report(window)
    data = context.reports.csv_bytes(report)
    assert isinstance(data, bytes)
    assert b"EMP-001" in data


def test_csv_exports_to_configured_folder(context, admin, team, frozen):
    report = context.reports.weekly_report(context.clock.week_bounds())
    filename = context.reports.suggest_filename(report, "csv")
    path = context.reports.write_csv(report, context.backups.exports_dir() / filename)
    assert path.parent.name == "exports"
    assert path.name.endswith(".csv")


def test_filename_suggestion_is_usable(context, admin, team, frozen):
    report = context.reports.daily_report(frozen.reference.date())
    name = context.reports.suggest_filename(report, "xlsx")
    assert name.endswith(".xlsx")
    assert " " not in name
    assert "daily" in name.lower() or "attendance" in name.lower()


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl not installed")
def test_xlsx_export(context, admin, team, frozen, tmp_path):
    from openpyxl import load_workbook

    report = context.reports.weekly_report(context.clock.week_bounds())
    path = context.reports.write_xlsx(report, tmp_path / "weekly.xlsx")
    assert path.is_file() and path.stat().st_size > 3000
    workbook = load_workbook(path)
    sheet = workbook.active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert any("EMP-001" in str(v) for v in values)
    assert any("Weekly Attendance Report" in str(v) for v in values)


@pytest.mark.skipif(not XLSX_AVAILABLE, reason="openpyxl not installed")
def test_xlsx_daily_export(context, admin, team, frozen, tmp_path):
    from openpyxl import load_workbook

    report = context.reports.daily_report(frozen.reference.date())
    path = context.reports.write_xlsx(report, tmp_path / "daily.xlsx")
    workbook = load_workbook(path)
    assert workbook.active.max_row > 5


def test_empty_report_still_exports(context, admin, tmp_path):
    report = context.reports.range_report(
        context.clock.today().replace(day=1), context.clock.today()
    )
    path = context.reports.write_csv(report, tmp_path / "empty.csv")
    assert path.is_file()
    assert context.reports.write_xlsx is not None